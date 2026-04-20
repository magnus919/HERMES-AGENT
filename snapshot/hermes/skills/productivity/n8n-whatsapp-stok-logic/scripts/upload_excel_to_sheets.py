#!/usr/bin/env python3
"""
Upload Excel inventory to Google Sheets using Service Account.
Run with: /usr/bin/python3 upload_excel_to_sheets.py

Prerequisites:
    pip install gspread google-auth openpyxl

Usage:
    1. Download Service Account JSON from Google Cloud Console
    2. Set path in SERVICE_ACCOUNT_FILE below
    3. Set EXCEL_FILE path
    4. Run script
"""

import gspread
from google.oauth2.service_account import Credentials
import openpyxl
import json
import os

# === CONFIGURATION ===
SERVICE_ACCOUNT_FILE = '/home/ubuntu/gcloud-service-account.json'
EXCEL_FILE = '/home/ubuntu/Daftar_Barang_Grosir.xlsx'
SPREADSHEET_NAME = 'Daftar Barang Grosir - Inventory'
SHEET_NAME = 'Stok Barang'
# ====================

SCOPES = [
    'https://www.googleapis.com/auth/spreadsheets',
    'https://www.googleapis.com/auth/drive.file'
]

def main():
    # Authenticate
    creds = Credentials.from_service_account_file(SERVICE_ACCOUNT_FILE, scopes=SCOPES)
    client = gspread.authorize(creds)

    # Read Excel
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    ws = wb.active

    # Get headers from row 1
    headers = []
    for cell in ws[1]:
        headers.append(str(cell.value) if cell.value else '')

    # Get data from row 2 onwards
    data = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1]:  # Skip rows where column B (Kode Barang) is empty
            data.append([str(cell) if cell else '' for cell in row])

    print(f"Headers ({len(headers)}): {headers}")
    print(f"Data rows: {len(data)}")

    # Create spreadsheet
    spreadsheet = client.create(SPREADSHEET_NAME)
    sheet = spreadsheet.sheet1
    sheet.title = SHEET_NAME

    # Update headers in row 1
    sheet.update('A1', [headers], value_input_option='RAW')

    # Update data starting from row 2
    if data:
        last_col = chr(65 + len(headers) - 1)  # Convert column count to letter (A, B, C, ...)
        range_end = f"{last_col}{len(data) + 1}"  # +1 for header row
        sheet.update(f'A2:{range_end}', data, value_input_option='RAW')

    # Make spreadsheet publicly readable
    spreadsheet.share('', perm_type='anyone', role='reader')

    # Save spreadsheet info
    info = {
        'url': spreadsheet.url,
        'id': spreadsheet.id,
        'sheet_name': SHEET_NAME,
        'rows': len(data) + 1,  # +1 for header
        'columns': len(headers)
    }

    output_file = '/home/ubuntu/google_sheets_info.json'
    with open(output_file, 'w') as f:
        json.dump(info, f, indent=2)

    print(f"\n✅ Success!")
    print(f"URL: {spreadsheet.url}")
    print(f"ID: {spreadsheet.id}")
    print(f"Info saved to: {output_file}")

if __name__ == '__main__':
    main()
